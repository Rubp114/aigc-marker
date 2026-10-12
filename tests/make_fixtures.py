"""Generate minimal test fixtures for xltx / dotx / potx / xmind / uof family / ofd and verify marking."""
import json
import os
import subprocess
import sys
import zipfile

FIXTURE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
MAIN_PY = r"D:\Code\aigc_marker\scripts\main.py"


def make_xltx(path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.template = True  # saved as template content type
    wb.active.title = "Sheet1"
    wb.active["A1"] = "hello"
    wb.save(path)


def _swap_main_content_type(src_path, dst_path, doc_ct, tpl_ct):
    """Rewrite [Content_Types].xml, replacing the main document content type."""
    with zipfile.ZipFile(src_path) as zin, \
         zipfile.ZipFile(dst_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(doc_ct.encode(), tpl_ct.encode())
            zout.writestr(item, data)


def make_dotx(path):
    import docx
    tmp = path + ".tmp.docx"
    doc = docx.Document()
    doc.add_paragraph("hello")
    doc.save(tmp)
    _swap_main_content_type(
        tmp, path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml")
    os.remove(tmp)


def make_potx(path):
    from pptx import Presentation
    tmp = path + ".tmp.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide
    prs.save(tmp)
    _swap_main_content_type(
        tmp, path,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml")
    os.remove(tmp)


def make_xmind_modern(path):
    content = [{
        "id": "sheet1", "class": "sheet", "title": "Sheet 1",
        "rootTopic": {"id": "root1", "class": "topic", "title": "Root Topic"},
    }]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("content.json", json.dumps(content, ensure_ascii=False))
        z.writestr("metadata.json", json.dumps({"creator": {"name": "test"}}))
        z.writestr("manifest.json", json.dumps({"file-entries": {}}))
        z.writestr("Thumbnails/thumbnail.png", b"\x89PNG fake")


def make_xmind_legacy(path):
    content_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0" version="2.0">'
        '<sheet id="s1"><title>Sheet 1</title>'
        '<topic id="t1"><title>Root Topic</title></topic>'
        '</sheet></xmap-content>')
    meta_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<meta xmlns="urn:xmind:xmap:xmlns:meta:2.0" version="2.0">'
        '<Creator><Name>test</Name></Creator></meta>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("content.xml", content_xml)
        z.writestr("meta.xml", meta_xml)
        z.writestr("META-INF/manifest.xml", "<manifest/>")


def make_uot(path):
    uot = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<uof:UOF xmlns:uof="http://schemas.uof.org/cn/2009/uof" '
        'xmlns:字="http://schemas.uof.org/cn/2009/wordproc">\n'
        '<uof:元数据><uof:标题>测试文档</uof:标题></uof:元数据>\n'
        '<字:文字处理文档>\n'
        '<字:分节>\n'
        '<字:段落 标识符="p1"><字:句><字:文本符>正文内容</字:文本符></字:句></字:段落>\n'
        '</字:分节>\n'
        '</字:文字处理文档>\n'
        '</uof:UOF>')
    with open(path, "w", encoding="utf-8") as f:
        f.write(uot)


def make_uos(path):
    uos = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<uof:UOF xmlns:uof="http://schemas.uof.org/cn/2009/uof" '
        'xmlns:表="http://schemas.uof.org/cn/2009/spreadsheet">\n'
        '<uof:元数据><uof:标题>测试表格</uof:标题></uof:元数据>\n'
        '<表:电子表格文档>\n'
        '<表:工作表集>\n'
        '<表:工作表 表:名称="Sheet1">\n'
        '<表:行><表:单元格><表:数据>原始数据</表:数据></表:单元格></表:行>\n'
        '</表:工作表>\n'
        '</表:工作表集>\n'
        '</表:电子表格文档>\n'
        '</uof:UOF>')
    with open(path, "w", encoding="utf-8") as f:
        f.write(uos)


def make_uop(path):
    uop = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<uof:UOF xmlns:uof="http://schemas.uof.org/cn/2009/uof" '
        'xmlns:演="http://schemas.uof.org/cn/2009/slideshow">\n'
        '<uof:元数据><uof:标题>测试演示</uof:标题></uof:元数据>\n'
        '<演:演示文稿文档>\n'
        '<演:幻灯片集>\n'
        '<演:幻灯片 标识符="s1">\n'
        '<演:段落><演:句><演:文本符>标题</演:文本符></演:句></演:段落>\n'
        '</演:幻灯片>\n'
        '</演:幻灯片集>\n'
        '</演:演示文稿文档>\n'
        '</uof:UOF>')
    with open(path, "w", encoding="utf-8") as f:
        f.write(uop)


def make_ofd(path):
    ofd_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016" DocType="OFD" Version="1.1">'
        '<ofd:DocBody>'
        '<ofd:DocInfo><ofd:DocID>abc</ofd:DocID><ofd:Title>test</ofd:Title></ofd:DocInfo>'
        '<ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot>'
        '</ofd:DocBody></ofd:OFD>')
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<ofd:Document xmlns:ofd="http://www.ofdspec.org/2016">'
        '<ofd:CommonData>'
        '<ofd:MaxUnitID>10</ofd:MaxUnitID>'
        '<ofd:PageArea><ofd:PhysicalBox>0 0 210 297</ofd:PhysicalBox></ofd:PageArea>'
        '<ofd:DocumentRes BaseLoc="Res/DocumentRes.xml"/>'
        '</ofd:CommonData>'
        '<ofd:Pages><ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/></ofd:Pages>'
        '</ofd:Document>')
    res_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<ofd:Res xmlns:ofd="http://www.ofdspec.org/2016">'
        '<ofd:Fonts><ofd:Font ID="3" FontName="SimSun"/></ofd:Fonts>'
        '</ofd:Res>')
    page_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<ofd:Page xmlns:ofd="http://www.ofdspec.org/2016">'
        '<ofd:Content><ofd:Layer ID="5">'
        '<ofd:TextObject ID="6" Boundary="10 10 80 10" Font="3" Size="5">'
        '<ofd:TextCode X="0" Y="4">Hello OFD</ofd:TextCode>'
        '</ofd:TextObject></ofd:Layer></ofd:Content>'
        '</ofd:Page>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("OFD.xml", ofd_xml)
        z.writestr("Doc_0/Document.xml", document_xml)
        z.writestr("Doc_0/Res/DocumentRes.xml", res_xml)
        z.writestr("Doc_0/Pages/Page_0/Content.xml", page_xml)


def run_main(path, extra=None):
    cmd = [sys.executable, MAIN_PY, path] + (extra or [])
    proc = subprocess.run(cmd, cwd=os.path.dirname(MAIN_PY),
                          capture_output=True, text=True, encoding="utf-8")
    return proc


def check(name, path, extra=None):
    print(f"\n=== {name} {' '.join(extra or [])} ===")
    p1 = run_main(path, extra)
    print("run1 rc:", p1.returncode)
    print(p1.stdout.strip())
    if p1.stderr.strip():
        print("STDERR:", p1.stderr.strip()[:500])
    # idempotency: second run should skip
    p2 = run_main(path, extra)
    skipped = "skipping" in p2.stdout
    print("run2 idempotent:", skipped)
    assert p1.returncode == 0, f"{name} first run failed"
    return path


def main():
    os.makedirs(FIXTURE_DIR, exist_ok=True)
    make_xltx(os.path.join(FIXTURE_DIR, "test.xltx"))
    make_dotx(os.path.join(FIXTURE_DIR, "test.dotx"))
    make_potx(os.path.join(FIXTURE_DIR, "test.potx"))
    make_xmind_modern(os.path.join(FIXTURE_DIR, "modern.xmind"))
    make_xmind_legacy(os.path.join(FIXTURE_DIR, "legacy.xmind"))
    make_uot(os.path.join(FIXTURE_DIR, "test.uot"))
    make_uot(os.path.join(FIXTURE_DIR, "test.uott"))  # same structure as uot
    make_uos(os.path.join(FIXTURE_DIR, "test.uos"))
    make_uos(os.path.join(FIXTURE_DIR, "test.uost"))  # same structure as uos
    make_uop(os.path.join(FIXTURE_DIR, "test.uop"))
    make_uop(os.path.join(FIXTURE_DIR, "test.uopt"))  # same structure as uop
    make_ofd(os.path.join(FIXTURE_DIR, "test.ofd"))

    check("xltx", os.path.join(FIXTURE_DIR, "test.xltx"))
    check("dotx", os.path.join(FIXTURE_DIR, "test.dotx"))
    check("potx", os.path.join(FIXTURE_DIR, "test.potx"))
    check("xmind-modern", os.path.join(FIXTURE_DIR, "modern.xmind"))
    check("xmind-legacy", os.path.join(FIXTURE_DIR, "legacy.xmind"))
    check("uot", os.path.join(FIXTURE_DIR, "test.uot"))
    check("uott", os.path.join(FIXTURE_DIR, "test.uott"))
    check("uos", os.path.join(FIXTURE_DIR, "test.uos"))
    check("uost", os.path.join(FIXTURE_DIR, "test.uost"))
    check("uop", os.path.join(FIXTURE_DIR, "test.uop"))
    check("uopt", os.path.join(FIXTURE_DIR, "test.uopt"))
    check("ofd", os.path.join(FIXTURE_DIR, "test.ofd"))
    check("ofd-skip-visible", os.path.join(FIXTURE_DIR, "test.ofd"), ["--skip-visible"])

    print("\nALL FIXTURE RUNS PASSED")


if __name__ == "__main__":
    main()
